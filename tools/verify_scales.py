#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Независимый аудит файла шкал: «ничего не потерялось?»

Сравнивает собранный docs/data_sample/scales.xlsx с исходным многовкладочным
docs/data_sample/legacy/scales_multisheet_2026.xlsx ЯЧЕЙКА ЗА ЯЧЕЙКОЙ:

  1. резервная копия побайтово равна оригиналу из git-истории;
  2. в новом файле ровно три листа: «Шкалы» (плоская таблица), «Витрина»
     (кросс-таблица на формулах) и «Блоки» (копия исходных вкладок);
  3. на листе «Блоки» 18 блоков-маркеров «### ДИЛЕР», порядок и имена совпадают
     с порядком вкладок исходника, скрытые вкладки помечены;
  4. для каждого блока: все значения, формулы (пересчитанные на новые строки),
     кэш-значения формул, форматы чисел, шрифты, заливки, границы, выравнивание,
     объединённые ячейки и высота строк перенесены 1:1, лишних ячеек нет;
  5. служебные «Лист1/Лист2/Лист3» в новый файл не переносились (по решению
     заказчика) и остались в резервной копии;
  6. плоская таблица: заголовки колонок, отсутствие объединённых ячеек и пустых
     строк (иначе сводная не строится), уникальность ключей, а главное — каждая
     запись ссылается на реальную ячейку исходника и значение в ней совпадает;
  7. витрина: все формулы смотрят в плоскую таблицу, а их кэш-значения равны
     тому, что должно получиться (INDEX/MATCH пересчитываются здесь же).

Логику разбора файла калькулятором проверяют JS-тесты (tests/*.test.mjs): они
сравнивают состояние приложения из плоской таблицы с состоянием из 18 вкладок.

Запуск:
    pip install openpyxl
    python3 tools/verify_scales.py            # полный аудит
    python3 tools/verify_scales.py --quiet    # только итог

Код возврата: 0 — всё сошлось, 1 — найдены расхождения.
"""

from __future__ import annotations

import argparse
import hashlib
import re
import subprocess
import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.formula.translate import Translator
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit('Нужен openpyxl:  pip install openpyxl')

ROOT = Path(__file__).resolve().parent.parent
NEW = ROOT / 'docs' / 'data_sample' / 'scales.xlsx'
LEGACY = ROOT / 'docs' / 'data_sample' / 'legacy' / 'scales_multisheet_2026.xlsx'
GIT_PATH_ORIGINAL = 'docs/data_sample/scales.xlsx'      # как файл лежал до сборки
GIT_PATH_BACKUP = 'docs/data_sample/legacy/scales_multisheet_2026.xlsx'

FLAT_SHEET = 'Шкалы'
SHOWCASE_SHEET = 'Витрина'
BLOCKS_SHEET = 'Блоки'

# Заголовки плоской таблицы — контракт между сборщиком, калькулятором и audit'ом.
FLAT_HEADER = [
    'Дилер', 'Тип', 'Раздел', 'Категория', 'Таблица', 'Показатель',
    '№ строки', 'Итог', 'Тир', '№ тира', 'Тир от', 'Тир до',
    'Значение', 'Единица', 'Вкладка (была)', 'Ячейка в исходнике',
]
FLAT_KEYS_HEADER = ['Ключ: строка', 'Ключ: тир', 'Ключ: значение',
                    'Ключ: порядок строки', 'Ключ: итог', 'Ключ: колонка плана']
LEGACY_SHEET_RE = re.compile(r'^\s*лист\s*\d*\s*$', re.I)
DEALER_MARKER_RE = re.compile(r'^\s*#{2,}\s*ДИЛЕР', re.I)
SKIP_MARKER_RE = re.compile(r'^\s*#{2,}\s*(РЕЕСТР|ВИТРИНА|ПЛОСКАЯ|СВОД|АРХИВ|СЛУЖЕБ)', re.I)
VYRUCHKA_RE = re.compile(r'Выручка\s*202[3-6]', re.I)
SISTEMA_RE = re.compile(r'СИСТЕМА\s+РАСЧЕТА\s+СКИДКИ\s+ПО\s+ЗАПЧАСТЯМ', re.I)

problems: list[str] = []
notes: list[str] = []
checked = {'cells': 0, 'styles': 0, 'formulas': 0, 'cached': 0, 'merges': 0,
           'heights': 0, 'flat': 0, 'flat_keys': 0, 'showcase': 0}


def fail(msg: str):
    problems.append(msg)


def note(msg: str):
    notes.append(msg)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


# ─── 1. Резервная копия == оригинал из git ──────────────────────────────────
def check_backup_matches_git():
    if not LEGACY.exists():
        fail(f'нет резервной копии: {LEGACY}')
        return
    try:
        commits = subprocess.run(
            ['git', 'log', '--all', '--format=%H', '--', GIT_PATH_ORIGINAL],
            cwd=ROOT, capture_output=True, text=True, check=True).stdout.split()
    except (subprocess.CalledProcessError, FileNotFoundError) as exc:
        note(f'git недоступен, сверка копии с историей пропущена ({exc})')
        return
    if not commits:
        note('в git-истории нет файла docs/data_sample/scales.xlsx — сверка пропущена')
        return
    sha = commits[-1]
    for path in (GIT_PATH_BACKUP, GIT_PATH_ORIGINAL):
        blob = subprocess.run(['git', 'show', f'{sha}:{path}'], cwd=ROOT,
                              capture_output=True)
        if blob.returncode == 0 and blob.stdout:
            git_hash = hashlib.sha256(blob.stdout).hexdigest()
            local_hash = sha256(LEGACY)
            if git_hash == local_hash:
                note(f'резервная копия побайтово равна оригиналу из git ({sha[:7]}:{path})')
            else:
                fail(f'резервная копия ОТЛИЧАЕТСЯ от оригинала в git ({sha[:7]}:{path})')
            return
    note('не удалось прочитать оригинал из git — сверка по хешу пропущена')


# ─── Сигнатура стиля ячейки (для сравнения оформления) ───────────────────────
def _color(x):
    if x is None:
        return None
    return (getattr(x, 'type', None), getattr(x, 'rgb', None), getattr(x, 'theme', None),
            round(getattr(x, 'tint', 0) or 0, 6), getattr(x, 'indexed', None))


def style_sig(cell):
    f, fill, b, al, p = cell.font, cell.fill, cell.border, cell.alignment, cell.protection
    sides = tuple((getattr(b, s).style, _color(getattr(b, s).color))
                  for s in ('left', 'right', 'top', 'bottom', 'diagonal'))
    return (
        cell.number_format,
        (f.name, f.sz, f.b, f.i, f.u, f.strike, f.vertAlign, _color(f.color)),
        (fill.fill_type, _color(fill.fgColor), _color(fill.bgColor)),
        sides,
        (al.horizontal, al.vertical, al.wrap_text, al.text_rotation, al.indent,
         al.shrink_to_fit),
        (p.locked, p.hidden),
    )


def value_equal(a, b):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool):
        return abs(float(a) - float(b)) <= 1e-9 * max(1.0, abs(float(a)))
    return a == b


def nonempty(ws):
    """Непустые ячейки листа: {(row, col): cell}."""
    out = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and not (isinstance(cell.value, str) and not cell.value.strip()):
                out[(cell.row, cell.column)] = cell
    return out


# ─── 2–3. Структура нового файла и блоки-маркеры ─────────────────────────────
def find_blocks(tgt):
    """Список блоков: (marker_row, marker_text, first_row, last_row)."""
    marks = []
    for r in range(1, tgt.max_row + 1):
        for c in range(1, min(4, tgt.max_column) + 1):
            v = tgt.cell(r, c).value
            if not isinstance(v, str):
                continue
            if DEALER_MARKER_RE.match(v):
                marks.append((r, v.strip(), 'dealer'))
            elif SKIP_MARKER_RE.match(v):
                marks.append((r, v.strip(), 'skip'))
            break
    blocks = []
    for i, (r, text, kind) in enumerate(marks):
        end = marks[i + 1][0] - 1 if i + 1 < len(marks) else tgt.max_row
        blocks.append({'marker_row': r, 'text': text, 'kind': kind,
                       'first': r + 1, 'last': end})
    return blocks


def block_info(tgt, block):
    """Имя дилера и тип по содержимому блока (независимо от текста маркера)."""
    header_row = None
    for r in range(block['first'], min(block['first'] + 8, block['last']) + 1):
        for c in range(1, tgt.max_column + 1):
            v = tgt.cell(r, c).value
            if isinstance(v, str) and VYRUCHKA_RE.search(v):
                header_row = r
                break
        if header_row:
            break
    if header_row is None:
        return None
    label_col, name = None, None
    for c in range(1, tgt.max_column + 1):
        v = tgt.cell(header_row + 1, c).value
        if isinstance(v, str) and v.strip():
            label_col, name = c, v.strip()
            break
    dtype = 'Заруб'
    if label_col:
        for r in range(block['first'], block['last'] + 1):
            v = tgt.cell(r, label_col).value
            if v and SISTEMA_RE.search(str(v)):
                dtype = 'РФ'
                break
    return {'header_row': header_row, 'label_col': label_col, 'name': name, 'type': dtype}


# ─── 4. Покадровое сравнение блока с исходной вкладкой ───────────────────────
def compare_block(src_ws, src_val_ws, tgt_ws, tgt_val_ws, row_shift, title):
    src_cells = nonempty(src_ws)
    tgt_cells = {}
    for (r, c), cell in nonempty(tgt_ws).items():
        if block_contains(r, row_shift, src_ws.max_row):
            tgt_cells[(r, c)] = cell

    # 4.1 значения / формулы / стили
    for (r, c), sc in src_cells.items():
        tr = r + row_shift
        tc = tgt_ws.cell(tr, c)
        checked['cells'] += 1
        sv, tv = sc.value, tc.value
        if isinstance(sv, str) and sv.startswith('='):
            checked['formulas'] += 1
            expected = Translator(sv, origin=sc.coordinate).translate_formula(tc.coordinate)
            if tv != expected:
                fail(f'{title}: формула {sc.coordinate} → {tc.coordinate}: '
                     f'ожидалось {expected!r}, получено {tv!r}')
            src_cached = src_val_ws.cell(r, c).value
            tgt_cached = tgt_val_ws.cell(tr, c).value
            checked['cached'] += 1
            if src_cached is not None and not value_equal(src_cached, tgt_cached):
                fail(f'{title}: кэш формулы {tc.coordinate}: было {src_cached!r}, '
                     f'стало {tgt_cached!r}')
        elif not value_equal(sv, tv):
            fail(f'{title}: значение {sc.coordinate} → {tc.coordinate}: '
                 f'было {sv!r}, стало {tv!r}')
        checked['styles'] += 1
        if sc.has_style or tc.has_style:
            s_sig, t_sig = style_sig(sc), style_sig(tc)
            if s_sig != t_sig:
                diff = [name for name, a, b in zip(
                    ('формат числа', 'шрифт', 'заливка', 'границы', 'выравнивание', 'защита'),
                    s_sig, t_sig) if a != b]
                fail(f'{title}: оформление {sc.coordinate} → {tc.coordinate} '
                     f'отличается ({", ".join(diff)})')
        tgt_cells.pop((tr, c), None)

    # 4.2 лишних ячеек в блоке быть не должно
    for (r, c), tc in sorted(tgt_cells.items()):
        fail(f'{title}: в блок попала ЛИШНЯЯ ячейка {tc.coordinate} = {tc.value!r}')

    # 4.3 объединённые ячейки
    src_merges = {str(m) for m in src_ws.merged_cells.ranges}
    tgt_merges = set()
    for m in tgt_ws.merged_cells.ranges:
        if block_contains(m.min_row, row_shift, src_ws.max_row) and \
                block_contains(m.max_row, row_shift, src_ws.max_row):
            shifted = f'{get_column_letter(m.min_col)}{m.min_row - row_shift}:' \
                      f'{get_column_letter(m.max_col)}{m.max_row - row_shift}'
            tgt_merges.add(shifted)
        checked['merges'] += 1
    missing, extra = src_merges - tgt_merges, tgt_merges - src_merges
    if missing:
        fail(f'{title}: потеряны объединённые ячейки {sorted(missing)}')
    if extra:
        fail(f'{title}: появились лишние объединённые ячейки {sorted(extra)}')

    # 4.4 высота строк
    for r, dim in src_ws.row_dimensions.items():
        if not dim.height:
            continue
        tgt_dim = tgt_ws.row_dimensions.get(r + row_shift)
        checked['heights'] += 1
        if tgt_dim is None or not tgt_dim.height or abs(tgt_dim.height - dim.height) > 0.01:
            fail(f'{title}: высота строки {r} ({dim.height}) не перенесена '
                 f'→ строка {r + row_shift} '
                 f'({tgt_dim.height if tgt_dim else None})')

    # 4.5 ширина колонок (берётся максимум по всем вкладкам, поэтому «не уже»)
    for letter, dim in src_ws.column_dimensions.items():
        if not dim.width:
            continue
        tgt_dim = tgt_ws.column_dimensions.get(letter)
        if tgt_dim is None or not tgt_dim.width or tgt_dim.width + 0.05 < dim.width:
            note(f'{title}: ширина колонки {letter} {dim.width} → '
                 f'{tgt_dim.width if tgt_dim else None} (на едином листе общая ширина)')


def block_contains(row, row_shift, src_max_row):
    return row_shift + 1 <= row <= row_shift + src_max_row


# ─── 6. Реестр ───────────────────────────────────────────────────────────────
def flat_columns(ws):
    """Колонки плоской таблицы по заголовкам + проверка самих заголовков."""
    header = [c.value for c in ws[1]]
    if header[:len(FLAT_HEADER)] != FLAT_HEADER:
        fail(f'заголовки листа «{FLAT_SHEET}» не совпадают с ожидаемыми: {header[:len(FLAT_HEADER)]}')
    keys = header[len(FLAT_HEADER):len(FLAT_HEADER) + len(FLAT_KEYS_HEADER)]
    if keys != FLAT_KEYS_HEADER:
        fail(f'служебные ключи листа «{FLAT_SHEET}» не совпадают: {keys}')
    return {name: i + 1 for i, name in enumerate(header) if name}


def check_flat(flat_ws, flat_val_ws, src_wb, src_val):
    """
    Плоская таблица (лист «Шкалы») — источник данных для калькулятора:
      • заголовки колонок — как договорились (по ним колонки ищет калькулятор);
      • нет объединённых ячеек и пустых строк: иначе сводная таблица не строится;
      • ключ значения уникален — по нему витрина ищет число через INDEX/MATCH;
      • каждая запись ссылается на реальную ячейку исходника, и значение в той
        ячейке совпадает с записанным (скидка нормализована в долю: 0.14 = 14%);
      • служебные ключи — формулы, и их кэш-значения совпадают с содержимым строки.
    """
    col = flat_columns(flat_ws)
    if len(col) != len(FLAT_HEADER) + len(FLAT_KEYS_HEADER):
        fail(f'в плоской таблице {len(col)} колонок, ожидалось '
             f'{len(FLAT_HEADER) + len(FLAT_KEYS_HEADER)}')

    if len(flat_ws.merged_cells.ranges):
        fail(f'в плоской таблице {len(flat_ws.merged_cells.ranges)} объединённых ячеек — '
             f'сводная таблица из такого листа не строится')
    if flat_ws.freeze_panes != 'B2':
        note(f'закрепление панелей листа «{FLAT_SHEET}»: {flat_ws.freeze_panes}')
    if not flat_ws.auto_filter.ref:
        fail('в плоской таблице нет автофильтра')

    dealers, keys, categories = {}, set(), set()
    n_scale = n_plan = 0
    last_row = flat_ws.max_row
    for r in range(2, last_row + 1):
        row = {name: flat_ws.cell(r, c).value for name, c in col.items()}
        if all(v is None or v == '' for v in row.values()):
            fail(f'«{FLAT_SHEET}»: пустая строка {r} внутри таблицы')
            continue
        dealer = row['Дилер']
        section = row['Раздел']
        category = row['Категория']
        indicator = row['Показатель']
        value = row['Значение']
        where = f'{dealer} / {section} / {category} / {indicator} / {row["Тир"]}'
        if not dealer or not section or not category or not indicator:
            fail(f'«{FLAT_SHEET}»!A{r}: у записи пустые обязательные колонки ({where})')
            continue
        dealers[dealer] = dealers.get(dealer, 0) + 1
        if not isinstance(value, (int, float)) or isinstance(value, bool):
            fail(f'«{FLAT_SHEET}»!M{r}: значение {value!r} не число ({where})')
            continue

        key = (dealer, section, category, row['№ строки'], row['№ тира'])
        if key in keys:
            fail(f'«{FLAT_SHEET}»: повтор ключа значения {"|".join(map(str, key))}')
        keys.add(key)

        # ── сверка с исходной ячейкой (главная проверка «ничего не потерялось») ──
        src_sheet, src_cell = row['Вкладка (была)'], row['Ячейка в исходнике']
        if section == 'Шкала':
            n_scale += 1
            categories.add(category)
            if not src_sheet or not src_cell:
                fail(f'«{FLAT_SHEET}»!A{r}: у значения шкалы нет ссылки на исходник ({where})')
            elif src_sheet not in src_wb.sheetnames:
                fail(f'«{FLAT_SHEET}»!A{r}: в исходнике нет вкладки «{src_sheet}»')
            else:
                raw = src_val[src_sheet][src_cell].value
                fmt = src_wb[src_sheet][src_cell].number_format
                expected = raw if '%' in (fmt or '') else (
                    raw / 100.0 if isinstance(raw, (int, float)) else None)
                checked['flat'] += 1
                if expected is None:
                    fail(f'«{FLAT_SHEET}»!M{r}: в исходнике {src_sheet}!{src_cell} пусто '
                         f'или не число ({raw!r}), а в таблице {value!r}')
                elif abs(expected - value) > 1e-9:
                    fail(f'«{FLAT_SHEET}»!M{r}: {value!r} ≠ исходник {src_sheet}!{src_cell} '
                         f'= {raw!r} ({expected!r} долей) — {where}')
                if row['Единица'] != '%':
                    fail(f'«{FLAT_SHEET}»!N{r}: единица значения шкалы {row["Единица"]!r} ({where})')
                if not 0 <= value <= 1:
                    fail(f'«{FLAT_SHEET}»!M{r}: скидка должна быть долей 0..1, а не {value!r} ({where})')
        else:
            n_plan += 1
            if src_sheet in src_wb.sheetnames and src_cell:
                raw = src_val[src_sheet][src_cell].value
                checked['flat'] += 1
                if not value_equal(value, raw):
                    fail(f'«{FLAT_SHEET}»!M{r}: {value!r} ≠ исходник {src_sheet}!{src_cell} = {raw!r}')

        # ── служебные ключи: формула + правильный кэш ──────────────────────
        expect_keys = {
            'Ключ: строка': f'{dealer}|{section}|{category}|{row["№ строки"]}',
            'Ключ: тир': f'{dealer}|{section}|{category}|{row["№ тира"]}',
            'Ключ: значение': f'{dealer}|{section}|{category}|{row["№ строки"]}|{row["№ тира"]}',
            'Ключ: порядок строки': f'{dealer}|{section}|{row["№ строки"]}',
            'Ключ: итог': (f'{dealer}|{section}|{category}|{row["№ тира"]}'
                           if row['Итог'] == 'да' else ''),
            'Ключ: колонка плана': f'{dealer}|{row["№ тира"]}' if section == 'План' else '',
        }
        for name, expected in expect_keys.items():
            c = col[name]
            formula = flat_ws.cell(r, c).value
            cached_value = flat_val_ws.cell(r, c).value
            if not isinstance(formula, str) or not formula.startswith('='):
                fail(f'«{FLAT_SHEET}»!{get_column_letter(c)}{r}: ключ «{name}» не формула')
                continue
            checked['flat_keys'] += 1
            if (cached_value or '') != expected:
                fail(f'«{FLAT_SHEET}»!{get_column_letter(c)}{r}: кэш ключа «{name}» = '
                     f'{cached_value!r}, ожидалось {expected!r}')

    if len(dealers) != 18:
        fail(f'в плоской таблице {len(dealers)} дилеров (ожидалось 18)')
    if n_scale < 800:
        fail(f'значений шкал всего {n_scale} — похоже, часть таблиц потерялась')
    if n_plan < 250:
        fail(f'плановых значений всего {n_plan} — похоже, шапка плана потерялась')
    if categories != {'Оборудование', 'Расходные материалы', 'Сервис (ЗЧ)'}:
        fail(f'категории шкал: {sorted(categories)}')
    note(f'плоская таблица: {len(keys)} записей ({n_scale} значений шкал, {n_plan} плановых) '
         f'по {len(dealers)} дилерам; сверено с исходником ячеек: {checked["flat"]}, '
         f'ключей: {checked["flat_keys"]}')
    return dealers


def showcase_layout(sc_val_ws):
    """Строки блоков витрины — ищем по заголовкам разделов в колонке A."""
    titles = {
        'plan': 'ПЛАН И БОНУСЫ',
        'equipment': 'ОБОРУДОВАНИЕ — шкала скидок',
        'materials': 'РАСХОДНЫЕ МАТЕРИАЛЫ — шкала скидок',
        'service': 'СЕРВИС (ЗЧ) — шкала скидок',
        'compare': 'СРАВНЕНИЕ ДИЛЕРОВ — строка ИТОГО по оборудованию',
    }
    found = {}
    for r in range(1, sc_val_ws.max_row + 1):
        v = sc_val_ws.cell(r, 1).value
        if isinstance(v, str):
            for key, title in titles.items():
                if v.strip() == title:
                    found[key] = r
    for key, title in titles.items():
        if key not in found:
            fail(f'на витрине нет раздела «{title}»')
    return found


def check_showcase(sc_ws, sc_val_ws, flat_ws, flat_val_ws, dealers):
    """
    Витрина: каждая формула смотрит в плоскую таблицу, а её кэш-значение равно
    тому, что должно получиться. INDEX/MATCH пересчитываем сами по плоской
    таблице — так ловятся и неверные ключи, и неверные ссылки, и пустой кэш.
    """
    col = {name: i + 1 for i, name in enumerate([c.value for c in flat_ws[1]]) if name}
    by_value, by_row, by_tier = {}, {}, {}
    by_order, by_total, by_plancol, type_of = {}, {}, {}, {}
    flat_dealers = []
    for r in range(2, flat_ws.max_row + 1):
        g = lambda name: flat_val_ws.cell(r, col[name]).value  # noqa: E731
        dealer, section, category = g('Дилер'), g('Раздел'), g('Категория')
        if not dealer:
            continue
        if dealer not in flat_dealers:
            flat_dealers.append(dealer)
        rec = {'dealer': dealer, 'section': section, 'category': category,
               'indicator': g('Показатель'), 'tier': g('Тир'), 'row_no': g('№ строки'),
               'tier_no': g('№ тира'), 'value': g('Значение'), 'is_total': g('Итог') == 'да',
               'table': g('Таблица')}
        by_value.setdefault(f'{dealer}|{section}|{category}|{rec["row_no"]}|{rec["tier_no"]}', rec)
        by_row.setdefault(f'{dealer}|{section}|{category}|{rec["row_no"]}', rec)
        by_tier.setdefault(f'{dealer}|{section}|{category}|{rec["tier_no"]}', rec)
        by_order.setdefault(f'{dealer}|{section}|{rec["row_no"]}', rec)
        if rec['is_total']:
            by_total.setdefault(f'{dealer}|{section}|{category}|{rec["tier_no"]}', rec)
        if section == 'План':
            by_plancol.setdefault(f'{dealer}|{rec["tier_no"]}', rec)
        type_of.setdefault(dealer, g('Тип'))

    d0 = sc_val_ws['B4'].value
    if d0 not in flat_dealers:
        fail(f'в B4 витрины выбран дилер {d0!r}, которого нет в плоской таблице')
        return
    layout = showcase_layout(sc_val_ws)
    if not layout:
        return

    # ── все формулы витрины ссылаются на плоскую таблицу и имеют кэш ────────
    formulas = cached = 0
    for row in sc_ws.iter_rows():
        for cell in row:
            v = cell.value
            if not isinstance(v, str) or not v.startswith('='):
                continue
            formulas += 1
            if 'INDEX' in v or 'MATCH' in v:
                if f"'{FLAT_SHEET}'!" not in v:
                    fail(f'«{SHOWCASE_SHEET}»!{cell.coordinate}: формула не смотрит в '
                         f'лист «{FLAT_SHEET}»: {v[:70]}')
            if sc_val_ws[cell.coordinate].value is not None:
                cached += 1
    if formulas == 0:
        fail('на витрине нет формул — она не будет обновляться при правке плоской таблицы')
    note(f'витрина: формул {formulas}, с кэш-значениями {cached} '
         f'(пустые — там, где у дилера нет данных)')

    def expect(coord, value, what):
        actual = sc_val_ws[coord].value
        checked['showcase'] += 1
        if value is None or value == '':
            if actual not in (None, ''):
                fail(f'«{SHOWCASE_SHEET}»!{coord}: {what} — ожидалось пусто, получено {actual!r}')
        elif isinstance(value, (int, float)) and not isinstance(value, bool):
            if not isinstance(actual, (int, float)) or abs(actual - value) > 1e-9:
                fail(f'«{SHOWCASE_SHEET}»!{coord}: {what} — ожидалось {value!r}, получено {actual!r}')
        elif str(actual or '') != str(value):
            fail(f'«{SHOWCASE_SHEET}»!{coord}: {what} — ожидалось {value!r}, получено {actual!r}')

    # ── шапка: тип, отчётность, заголовок сервиса ─────────────────────────
    expect('F4', type_of.get(d0, ''), 'тип дилера')
    expect('H4', 'за квартал' if type_of.get(d0) == 'РФ' else 'за полугодие', 'отчётность')
    svc = by_value.get(f'{d0}|Шкала|Сервис (ЗЧ)|1|1')
    expect('J4', svc['table'] if svc else '', 'заголовок шкалы сервиса')

    def table(title_row, section, category, max_rows, label_field):
        """Проверяет таблицу витрины: подписи, тиры и все значения."""
        hdr = title_row + 1
        for j in range(1, 13):
            coord = f'{get_column_letter(1 + j)}{hdr}'
            if section == 'План':
                rec = by_plancol.get(f'{d0}|{j}')
            else:
                rec = by_tier.get(f'{d0}|{section}|{category}|{j}')
            field = 'indicator' if section == 'План' else 'tier'
            expect(coord, rec[field] if rec else '', f'подпись колонки {j}')
        for k in range(1, max_rows + 1):
            r = hdr + k
            if section == 'План':
                lab = by_order.get(f'{d0}|План|{k}')
                label = lab['category'] if lab else ''
            else:
                lab = by_row.get(f'{d0}|Шкала|{category}|{k}')
                label = lab['indicator'] if lab else ''
            expect(f'A{r}', label, f'подпись строки {k}')
            for j in range(1, 13):
                coord = f'{get_column_letter(1 + j)}{r}'
                if sc_ws[coord].value is None:
                    continue
                if section == 'План':
                    rec = by_value.get(f'{d0}|План|{label}|{k}|{j}')
                else:
                    rec = by_value.get(f'{d0}|Шкала|{category}|{k}|{j}')
                expect(coord, rec['value'] if rec else '', f'{label} / колонка {j}')

    table(layout['plan'], 'План', 'Итого по дилеру', 5, 'category')
    table(layout['equipment'], 'Шкала', 'Оборудование', 9, 'indicator')
    table(layout['materials'], 'Шкала', 'Расходные материалы', 4, 'indicator')
    table(layout['service'], 'Шкала', 'Сервис (ЗЧ)', 1, 'indicator')

    # ── сравнение дилеров: строка ИТОГО по оборудованию ───────────────────
    hdr = layout['compare'] + 1
    for j in range(1, 7):
        rec = by_tier.get(f'{d0}|Шкала|Оборудование|{j}')
        expect(f'{get_column_letter(1 + j)}{hdr}', rec['tier'] if rec else '',
               f'тир {j} в сравнении')
    for i, dealer in enumerate(flat_dealers):
        r = hdr + 1 + i
        expect(f'A{r}', dealer, f'дилер {i + 1} в сравнении')
        expect(f'H{r}', type_of.get(dealer, ''), f'тип {dealer}')
        for j in range(1, 7):
            rec = by_total.get(f'{dealer}|Шкала|Оборудование|{j}')
            expect(f'{get_column_letter(1 + j)}{r}', rec['value'] if rec else '',
                   f'ИТОГО {dealer} / тир {j}')

    # ── выпадающий список дилеров ─────────────────────────────────────────
    dvs = [dv for dv in sc_ws.data_validations.dataValidation
           if 'B4' in str(dv.sqref)]
    if not dvs:
        fail('в B4 витрины нет выпадающего списка дилеров')
    else:
        m = re.search(r'\$W\$(\d+):\$W\$(\d+)', dvs[0].formula1 or '')
        if not m:
            fail(f'непонятный источник списка дилеров: {dvs[0].formula1!r}')
        else:
            listed = [sc_val_ws.cell(r, 23).value
                      for r in range(int(m.group(1)), int(m.group(2)) + 1)]
            if [x for x in listed if x] != flat_dealers:
                fail('список дилеров в витрине не совпадает с плоской таблицей')
            note(f'витрина: выпадающий список из {len(flat_dealers)} дилеров, '
                 f'сверено значений {checked["showcase"]}')



def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--new', type=Path, default=NEW)
    ap.add_argument('--legacy', type=Path, default=LEGACY)
    ap.add_argument('--quiet', action='store_true')
    args = ap.parse_args()

    if not args.new.exists():
        sys.exit(f'нет файла {args.new} — сначала запустите tools/consolidate_scales.py')
    if not args.legacy.exists():
        sys.exit(f'нет резервной копии {args.legacy}')

    print('Аудит файла шкал (плоская таблица + витрина + блоки)')
    print('=' * 78)
    check_backup_matches_git()

    src_wb = openpyxl.load_workbook(args.legacy, data_only=False)
    src_val = openpyxl.load_workbook(args.legacy, data_only=True)
    new_wb = openpyxl.load_workbook(args.new, data_only=False)
    new_val = openpyxl.load_workbook(args.new, data_only=True)

    # 2. структура
    expected = [FLAT_SHEET, SHOWCASE_SHEET, BLOCKS_SHEET]
    if new_wb.sheetnames != expected:
        fail(f'ожидались листы {expected}, получено {new_wb.sheetnames}')
    tgt_ws, tgt_val_ws = new_wb[BLOCKS_SHEET], new_val[BLOCKS_SHEET]
    flat_ws, flat_val_ws = new_wb[FLAT_SHEET], new_val[FLAT_SHEET]

    # 3. блоки
    blocks = find_blocks(tgt_ws)
    dealer_blocks = [b for b in blocks if b['kind'] == 'dealer']
    src_dealer_sheets = [ws for ws in src_wb.worksheets if not LEGACY_SHEET_RE.match(ws.title)]
    src_service_sheets = [ws.title for ws in src_wb.worksheets if LEGACY_SHEET_RE.match(ws.title)]
    if len(dealer_blocks) != len(src_dealer_sheets):
        fail(f'блоков «### ДИЛЕР» {len(dealer_blocks)}, дилерских вкладок в исходнике '
             f'{len(src_dealer_sheets)}')

    blocks_geometry = []
    for i, (block, src_ws) in enumerate(zip(dealer_blocks, src_dealer_sheets), 1):
        title = f'блок {i} ({src_ws.title})'
        info = block_info(tgt_ws, block)
        if info is None:
            fail(f'{title}: в блоке не найдена шапка «Выручка …»')
            continue
        if info['name'] is None:
            fail(f'{title}: не найдено имя дилера')
            continue
        # имя и тип в маркере должны совпадать с содержимым блока
        for what, expected in (('имя', info['name']), ('тип', info['type'])):
            if expected not in block['text']:
                fail(f'{title}: в маркере нет {what} {expected!r}: {block["text"][:80]}')
        if src_ws.title not in block['text']:
            fail(f'{title}: в маркере нет имени исходной вкладки «{src_ws.title}»')
        if src_ws.sheet_state != 'visible' and 'скрыт' not in block['text'].lower():
            fail(f'{title}: исходная вкладка была скрыта, но в маркере нет пометки')
        # блок должен начинаться сразу за маркером и содержать все строки вкладки
        if block['first'] != block['marker_row'] + 1:
            fail(f'{title}: блок начинается не сразу за маркером')
        row_shift = block['marker_row']
        compare_block(src_ws, src_val[src_ws.title], tgt_ws, tgt_val_ws, row_shift, title)
        blocks_geometry.append({
            'title': src_ws.title, 'first': block['first'],
            'last': row_shift + src_ws.max_row, 'label_col': info['label_col'],
        })

    # 5. служебные листы
    for name in src_service_sheets:
        if name in new_wb.sheetnames:
            note(f'служебный лист «{name}» присутствует в новом файле')
        cells = len(nonempty(src_wb[name]))
        note(f'служебный лист «{name}» ({cells} непустых ячеек) в новый файл не переносился — '
             f'остался в резервной копии {args.legacy.name}')
    if len(tgt_ws.merged_cells.ranges) != sum(
            len(ws.merged_cells.ranges) for ws in src_dealer_sheets):
        fail(f'число объединённых ячеек на листе «{BLOCKS_SHEET}» не равно сумме по вкладкам')

    # 6. плоская таблица (источник данных) и 7. витрина (представление)
    dealers = check_flat(flat_ws, flat_val_ws, src_wb, src_val)
    check_showcase(new_wb[SHOWCASE_SHEET], new_val[SHOWCASE_SHEET],
                   flat_ws, flat_val_ws, dealers)

    print('\nПроверено:')
    for key, label in (('cells', 'ячеек блоков (значение/формула)'), ('formulas', 'формул блоков'),
                       ('cached', 'кэш-значений формул'), ('styles', 'ячеек (оформление)'),
                       ('merges', 'объединённых ячеек'), ('heights', 'высот строк'),
                       ('flat', 'значений плоской таблицы'), ('flat_keys', 'служебных ключей'),
                       ('showcase', 'ячеек витрины')):
        print(f'  {label:<28}: {checked[key]}')

    if notes and not args.quiet:
        print('\nЗаметки:')
        for n in notes:
            print(f'  • {n}')

    print()
    if problems:
        print(f'✗ НАЙДЕНО РАСХОЖДЕНИЙ: {len(problems)}')
        for p in problems[:60]:
            print(f'  - {p}')
        if len(problems) > 60:
            print(f'  … и ещё {len(problems) - 60}')
        return 1
    print('✓ Всё сошлось: плоская таблица, витрина и блоки содержат те же данные, '
          'что и 18 вкладок исходника.')
    return 0


if __name__ == '__main__':
    sys.exit(main())
